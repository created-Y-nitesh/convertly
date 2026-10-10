import io
import json
import pytest
from pathlib import Path
from unittest.mock import patch, MagicMock
from fastapi.testclient import TestClient
import fitz
import docx
import openpyxl
from PIL import Image

from app.main import app
from app.services.storage_service import storage_service
from app.core.config import settings
from app.db.mongodb import connect_to_mongodb

client = TestClient(app)

def create_sample_pdf() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text(fitz.Point(50, 72), "Convertly PDF Test Document\nHello world from PyMuPDF", fontsize=14)
    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    return buf.getvalue()

def create_sample_docx() -> bytes:
    d = docx.Document()
    d.add_heading("Convertly Word Test", level=1)
    d.add_paragraph("Paragraph 1: Testing docx conversion pipeline.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()

def create_sample_png() -> bytes:
    img = Image.new("RGBA", (100, 100), color=(255, 0, 0, 255))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()

def create_sample_jpg() -> bytes:
    img = Image.new("RGB", (100, 100), color=(0, 255, 0))
    buf = io.BytesIO()
    img.save(buf, format="JPEG")
    return buf.getvalue()

def create_sample_webp() -> bytes:
    img = Image.new("RGBA", (100, 100), color=(0, 0, 255, 255))
    buf = io.BytesIO()
    img.save(buf, format="WEBP")
    return buf.getvalue()

def create_sample_csv() -> bytes:
    content = "Name,Department,Salary\nAlice,Engineering,120000\nBob,Design,95000\nCharlie,Marketing,88000\n"
    return content.encode("utf-8")

def create_sample_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Employees"
    ws.append(["ID", "Name", "Score"])
    ws.append([1, "Alpha", 98])
    ws.append([2, "Beta", 85])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def create_sample_txt() -> bytes:
    return b"Plain text file content\nLine 2 of document\nLine 3 of document\n"

# 1. PDF -> DOCX
def test_convert_pdf_to_docx():
    pdf_bytes = create_sample_pdf()
    res = client.post(
        "/api/v1/conversions",
        data={"source_format": "pdf", "target_format": "docx"},
        files={"file": ("sample.pdf", pdf_bytes, "application/pdf")}
    )
    assert res.status_code == 200, res.text
    data = res.json()
    assert data["status"] == "completed"
    assert data["result_file_size"] > 0
    assert "token=" in data["download_url"]
    assert data["access_token"] is not None

    # Test authorized download endpoint via download_url
    dl_res = client.get(data["download_url"])
    assert dl_res.status_code == 200
    assert len(dl_res.content) == data["result_file_size"]

    # Verify output is valid docx
    out_docx = docx.Document(io.BytesIO(dl_res.content))
    paragraphs_text = " ".join(p.text for p in out_docx.paragraphs)
    assert "Convertly PDF Test Document" in paragraphs_text

# 2. PDF -> PNG
def test_convert_pdf_to_png():
    pdf_bytes = create_sample_pdf()
    res = client.post(
        "/api/v1/conversions",
        data={"source_format": "pdf", "target_format": "png"},
        files={"file": ("sample.pdf", pdf_bytes, "application/pdf")}
    )
    assert res.status_code == 200, res.text
    data = res.json()

    dl_res = client.get(data["download_url"])
    assert dl_res.status_code == 200
    img = Image.open(io.BytesIO(dl_res.content))
    assert img.format == "PNG"
    assert img.width > 0 and img.height > 0

# 3. PDF -> JPG
def test_convert_pdf_to_jpg():
    pdf_bytes = create_sample_pdf()
    res = client.post(
        "/api/v1/conversions",
        data={"source_format": "pdf", "target_format": "jpg"},
        files={"file": ("sample.pdf", pdf_bytes, "application/pdf")}
    )
    assert res.status_code == 200, res.text
    data = res.json()

    dl_res = client.get(data["download_url"])
    assert dl_res.status_code == 200
    img = Image.open(io.BytesIO(dl_res.content))
    assert img.format == "JPEG"

# 4. PDF -> TXT
def test_convert_pdf_to_txt():
    pdf_bytes = create_sample_pdf()
    res = client.post(
        "/api/v1/conversions",
        data={"source_format": "pdf", "target_format": "txt"},
        files={"file": ("sample.pdf", pdf_bytes, "application/pdf")}
    )
    assert res.status_code == 200, res.text
    data = res.json()

    dl_res = client.get(data["download_url"])
    assert dl_res.status_code == 200
    text_content = dl_res.content.decode("utf-8")
    assert "Convertly PDF Test Document" in text_content

# 5. DOCX -> TXT
def test_convert_docx_to_txt():
    docx_bytes = create_sample_docx()
    res = client.post(
        "/api/v1/conversions",
        data={"source_format": "docx", "target_format": "txt"},
        files={"file": ("sample.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert res.status_code == 200, res.text
    data = res.json()

    dl_res = client.get(data["download_url"])
    assert dl_res.status_code == 200
    text_content = dl_res.content.decode("utf-8")
    assert "Testing docx conversion pipeline" in text_content

# 6. TXT -> PDF
def test_convert_txt_to_pdf():
    txt_bytes = create_sample_txt()
    res = client.post(
        "/api/v1/conversions",
        data={"source_format": "txt", "target_format": "pdf"},
        files={"file": ("sample.txt", txt_bytes, "text/plain")}
    )
    assert res.status_code == 200, res.text
    data = res.json()

    dl_res = client.get(data["download_url"])
    assert dl_res.status_code == 200
    doc = fitz.open(stream=dl_res.content, filetype="pdf")
    assert doc.page_count >= 1
    extracted = doc[0].get_text()
    assert "Plain text file content" in extracted

# 7. CSV -> XLSX
def test_convert_csv_to_xlsx():
    csv_bytes = create_sample_csv()
    res = client.post(
        "/api/v1/conversions",
        data={"source_format": "csv", "target_format": "xlsx"},
        files={"file": ("records.csv", csv_bytes, "text/csv")}
    )
    assert res.status_code == 200, res.text
    data = res.json()

    dl_res = client.get(data["download_url"])
    assert dl_res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(dl_res.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Name", "Department", "Salary")
    assert rows[1] == ("Alice", "Engineering", "120000")

# 8. XLSX -> CSV
def test_convert_xlsx_to_csv():
    xlsx_bytes = create_sample_xlsx()
    res = client.post(
        "/api/v1/conversions",
        data={"source_format": "xlsx", "target_format": "csv"},
        files={"file": ("data.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert res.status_code == 200, res.text
    data = res.json()

    dl_res = client.get(data["download_url"])
    assert dl_res.status_code == 200
    lines = dl_res.content.decode("utf-8").strip().splitlines()
    assert "ID,Name,Score" in lines[0]
    assert "1,Alpha,98" in lines[1]

# 9. PNG -> JPG
def test_convert_png_to_jpg():
    png_bytes = create_sample_png()
    res = client.post(
        "/api/v1/conversions",
        data={"source_format": "png", "target_format": "jpg"},
        files={"file": ("image.png", png_bytes, "image/png")}
    )
    assert res.status_code == 200, res.text
    data = res.json()

    dl_res = client.get(data["download_url"])
    assert dl_res.status_code == 200
    img = Image.open(io.BytesIO(dl_res.content))
    assert img.format == "JPEG"

# 10. JPG -> PNG
def test_convert_jpg_to_png():
    jpg_bytes = create_sample_jpg()
    res = client.post(
        "/api/v1/conversions",
        data={"source_format": "jpg", "target_format": "png"},
        files={"file": ("image.jpg", jpg_bytes, "image/jpeg")}
    )
    assert res.status_code == 200, res.text
    data = res.json()

    dl_res = client.get(data["download_url"])
    assert dl_res.status_code == 200
    img = Image.open(io.BytesIO(dl_res.content))
    assert img.format == "PNG"

# 11. PNG -> WEBP
def test_convert_png_to_webp():
    png_bytes = create_sample_png()
    res = client.post(
        "/api/v1/conversions",
        data={"source_format": "png", "target_format": "webp"},
        files={"file": ("image.png", png_bytes, "image/png")}
    )
    assert res.status_code == 200, res.text
    data = res.json()

    dl_res = client.get(data["download_url"])
    assert dl_res.status_code == 200
    img = Image.open(io.BytesIO(dl_res.content))
    assert img.format == "WEBP"

# 12. WEBP -> PNG
def test_convert_webp_to_png():
    webp_bytes = create_sample_webp()
    res = client.post(
        "/api/v1/conversions",
        data={"source_format": "webp", "target_format": "png"},
        files={"file": ("image.webp", webp_bytes, "image/webp")}
    )
    assert res.status_code == 200, res.text
    data = res.json()

    dl_res = client.get(data["download_url"])
    assert dl_res.status_code == 200
    img = Image.open(io.BytesIO(dl_res.content))
    assert img.format == "PNG"

# 13. PNG -> PDF
def test_convert_png_to_pdf():
    png_bytes = create_sample_png()
    res = client.post(
        "/api/v1/conversions",
        data={"source_format": "png", "target_format": "pdf"},
        files={"file": ("image.png", png_bytes, "image/png")}
    )
    assert res.status_code == 200, res.text
    data = res.json()

    dl_res = client.get(data["download_url"])
    assert dl_res.status_code == 200
    doc = fitz.open(stream=dl_res.content, filetype="pdf")
    assert doc.page_count == 1

# 14. JPG -> PDF
def test_convert_jpg_to_pdf():
    jpg_bytes = create_sample_jpg()
    res = client.post(
        "/api/v1/conversions",
        data={"source_format": "jpg", "target_format": "pdf"},
        files={"file": ("image.jpg", jpg_bytes, "image/jpeg")}
    )
    assert res.status_code == 200, res.text
    data = res.json()

    dl_res = client.get(data["download_url"])
    assert dl_res.status_code == 200
    doc = fitz.open(stream=dl_res.content, filetype="pdf")
    assert doc.page_count == 1

# 15. Unsupported route rejection
def test_unsupported_conversion_route():
    res = client.post(
        "/api/v1/conversions",
        data={"source_format": "csv", "target_format": "png"},
        files={"file": ("test.csv", b"a,b,c", "text/csv")}
    )
    assert res.status_code == 400
    assert "not supported" in res.json()["detail"].lower()

# 16. Corrupted file error handling
def test_corrupted_pdf_handling():
    corrupted_bytes = b"%PDF-1.4\ncorrupted content and random garbage byte sequence"
    res = client.post(
        "/api/v1/conversions",
        data={"source_format": "pdf", "target_format": "docx"},
        files={"file": ("broken.pdf", corrupted_bytes, "application/pdf")}
    )
    assert res.status_code in [400, 500]

# 17. Empty file handling
def test_empty_file_handling():
    res = client.post(
        "/api/v1/conversions",
        data={"source_format": "pdf", "target_format": "docx"},
        files={"file": ("empty.pdf", b"", "application/pdf")}
    )
    assert res.status_code == 400
    assert "empty" in res.json()["detail"].lower()

# 18. Capability Access Token Enforcement on Download
def test_download_capability_token_enforcement():
    txt_bytes = b"Capability token test document"
    res = client.post(
        "/api/v1/conversions",
        data={"source_format": "txt", "target_format": "pdf"},
        files={"file": ("token_test.txt", txt_bytes, "text/plain")}
    )
    assert res.status_code == 200
    job_id = res.json()["job_id"]
    access_token = res.json()["access_token"]
    assert access_token is not None

    # A) Attempt download with NO token -> 403 Forbidden
    no_token_res = client.get(f"/api/v1/conversions/{job_id}/download")
    assert no_token_res.status_code == 403
    assert "capability token" in no_token_res.json()["detail"].lower()

    # B) Attempt download with FAKE spoofed token -> 403 Forbidden
    fake_token_res = client.get(f"/api/v1/conversions/{job_id}/download?token=fake_spoofed_token_12345")
    assert fake_token_res.status_code == 403

    # C) Authorized download with valid query token -> 200 OK
    auth_query_res = client.get(f"/api/v1/conversions/{job_id}/download?token={access_token}")
    assert auth_query_res.status_code == 200
    assert auth_query_res.content.startswith(b"%PDF")

    # D) Authorized download with Bearer Header -> 200 OK
    auth_header_res = client.get(
        f"/api/v1/conversions/{job_id}/download",
        headers={"Authorization": f"Bearer {access_token}"}
    )
    assert auth_header_res.status_code == 200

# 19. Cleanup verification and token-protected deletion
def test_job_cleanup_and_deletion_authorization():
    txt_bytes = b"Temp cleanup test"
    res = client.post(
        "/api/v1/conversions",
        data={"source_format": "txt", "target_format": "pdf"},
        files={"file": ("temp.txt", txt_bytes, "text/plain")}
    )
    assert res.status_code == 200
    job_id = res.json()["job_id"]
    token = res.json()["access_token"]

    # Delete without token -> 403 Forbidden
    unauth_del = client.delete(f"/api/v1/conversions/{job_id}")
    assert unauth_del.status_code == 403

    # Delete with valid token -> 200 OK
    del_res = client.delete(f"/api/v1/conversions/{job_id}?token={token}")
    assert del_res.status_code == 200

    # Clean up storage
    cleanup_res = storage_service.cleanup_job(job_id)
    assert cleanup_res is True
    job_dir = storage_service.base_dir / job_id
    assert not job_dir.exists()

# 20. Signature mismatch strict rejection
def test_signature_mismatch_rejected():
    fake_pdf = b"THIS_IS_NOT_A_VALID_PDF_HEADER_123456789"
    res = client.post(
        "/api/v1/conversions",
        data={"source_format": "pdf", "target_format": "docx"},
        files={"file": ("fake.pdf", fake_pdf, "application/pdf")}
    )
    assert res.status_code == 400
    assert "signature mismatch" in res.json()["detail"].lower()

# 21. DOCX-to-PDF Engine Availability or Explicit Error Test
def test_convert_docx_to_pdf_behavior():
    docx_bytes = create_sample_docx()
    res = client.post(
        "/api/v1/conversions",
        data={"source_format": "docx", "target_format": "pdf"},
        files={"file": ("sample.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    # If LibreOffice/Gotenberg is present on system -> 200 with valid PDF
    # If neither is installed -> 503 with explicit engine guidance
    if res.status_code == 200:
        pdf_content = client.get(res.json()["download_url"]).content
        doc = fitz.open(stream=pdf_content, filetype="pdf")
        assert doc.page_count >= 1
    else:
        assert res.status_code == 503
        assert "libreoffice" in res.json()["detail"].lower() or "gotenberg" in res.json()["detail"].lower()

# 22. Mocked Full Execution of DOCX-to-PDF Engine
def test_mocked_docx_to_pdf_execution():
    docx_bytes = create_sample_docx()
    
    # Mock find_libreoffice_binary and subprocess.run to simulate LibreOffice container environment
    with patch("app.converters.office_converter.find_libreoffice_binary", return_value="/usr/bin/soffice"):
        def fake_run(cmd, *args, **kwargs):
            outdir = cmd[cmd.index("--outdir") + 1]
            infile = cmd[-1]
            out_pdf = Path(outdir) / f"{Path(infile).stem}.pdf"
            # Write valid PDF
            doc = fitz.open()
            page = doc.new_page()
            page.insert_text(fitz.Point(50, 72), "Converted from DOCX via LibreOffice", fontsize=12)
            doc.save(str(out_pdf))
            doc.close()
            mock_proc = MagicMock()
            mock_proc.returncode = 0
            return mock_proc

        with patch("subprocess.run", side_effect=fake_run):
            res = client.post(
                "/api/v1/conversions",
                data={"source_format": "docx", "target_format": "pdf"},
                files={"file": ("document.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
            )
            assert res.status_code == 200
            dl_res = client.get(res.json()["download_url"])
            assert dl_res.status_code == 200
            pdf_doc = fitz.open(stream=dl_res.content, filetype="pdf")
            assert "Converted from DOCX via LibreOffice" in pdf_doc[0].get_text()

# 23. Google Sheets Integration Ping Endpoint Test
def test_google_sheets_ping_endpoint():
    # Tests the server-side test endpoint without exposing secrets
    res = client.post("/api/v1/integrations/google-sheets/test")
    assert res.status_code == 200
    data = res.json()
    assert "success" in data
    assert "message" in data

# 24. Anonymous History Redaction Test
def test_history_anonymous_redaction():
    res = client.get("/api/v1/history")
    assert res.status_code == 200
    items = res.json()
    if items:
        # Without X-Admin-Key, filename must be masked
        first_item = items[0]
        assert "document_" in first_item["original_filename"]

# 25. Production Persistence Startup Validation
@pytest.mark.asyncio
async def test_production_database_strict_mode():
    with patch.object(settings, "ENVIRONMENT", "production"):
        with patch.object(settings, "ALLOW_IN_MEMORY_FALLBACK", False):
            with patch("motor.motor_asyncio.AsyncIOMotorClient.admin", create=True) as mock_admin:
                mock_admin.command.side_effect = Exception("Connection refused to Atlas")
                with pytest.raises(RuntimeError) as exc_info:
                    await connect_to_mongodb()
                assert "FATAL: Production database connection failed" in str(exc_info.value)

# 26. Storage purge method test
def test_storage_purge_expired():
    purged = storage_service.purge_expired_files(max_age_seconds=0)
    assert isinstance(purged, int)
