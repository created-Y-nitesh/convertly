import csv
from pathlib import Path
from typing import Dict, Any
import openpyxl

def convert_csv_to_xlsx(input_path: Path, output_path: Path, options: Dict[str, Any]) -> None:
    delimiter = options.get("csv_delimiter", ",")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    
    with open(input_path, "r", encoding="utf-8", errors="replace") as f:
        reader = csv.reader(f, delimiter=delimiter)
        for row in reader:
            ws.append(row)
            
    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 50)
        
    wb.save(output_path)

def convert_xlsx_to_csv(input_path: Path, output_path: Path, options: Dict[str, Any]) -> None:
    delimiter = options.get("csv_delimiter", ",")
    wb = openpyxl.load_workbook(input_path, data_only=True)
    ws = wb.active
    
    with open(output_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f, delimiter=delimiter)
        for row in ws.iter_rows(values_only=True):
            writer.writerow([str(cell) if cell is not None else "" for cell in row])
