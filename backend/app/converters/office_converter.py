import os
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Dict, Any, Optional
import docx
import httpx
from fastapi import HTTPException

from app.core.config import settings
from app.core.logging import logger

def find_libreoffice_binary() -> Optional[str]:
    """
    Locates the soffice/libreoffice executable across PATH and known OS directories.
    """
    # 1. Custom path override
    if getattr(settings, "LIBREOFFICE_PATH", "") and os.path.isfile(settings.LIBREOFFICE_PATH):
        return settings.LIBREOFFICE_PATH

    # 2. System PATH
    for name in ["soffice", "libreoffice"]:
        found = shutil.which(name)
        if found:
            return found

    # 3. Standard Windows installation paths
    win_paths = [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    for p in win_paths:
        if os.path.isfile(p):
            return p

    # 4. Standard Linux / Docker paths
    linux_paths = [
        "/usr/bin/soffice",
        "/usr/bin/libreoffice",
        "/usr/local/bin/soffice",
    ]
    for p in linux_paths:
        if os.path.isfile(p):
            return p

    return None

def is_docx_to_pdf_engine_available() -> bool:
    """
    Returns True if either a local LibreOffice binary or Gotenberg URL is available.
    """
    return find_libreoffice_binary() is not None or bool(getattr(settings, "GOTENBERG_URL", ""))

def convert_docx_to_pdf(input_path: Path, output_path: Path, options: Dict[str, Any]) -> None:
    """
    Executes high-fidelity DOCX to PDF conversion via LibreOffice headless or Gotenberg service.
    Preserves document styles, tables, vector imagery, and page layouts.
    Raises HTTP 503 if no genuine engine is installed in the current environment.
    """
    # 1. Gotenberg microservice route (if configured)
    gotenberg_url = getattr(settings, "GOTENBERG_URL", "").strip()
    if gotenberg_url:
        try:
            endpoint = f"{gotenberg_url.rstrip('/')}/forms/libreoffice/convert"
            with open(input_path, "rb") as f:
                files = {"files": (input_path.name, f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
                res = httpx.post(endpoint, files=files, timeout=60.0)
            if res.status_code == 200 and len(res.content) > 0:
                with open(output_path, "wb") as out_f:
                    out_f.write(res.content)
                return
            else:
                logger.error(f"Gotenberg DOCX-to-PDF error: HTTP {res.status_code} - {res.text[:200]}")
                raise RuntimeError(f"Gotenberg conversion service failed with HTTP {res.status_code}")
        except Exception as e:
            logger.error(f"Gotenberg connection error: {str(e)}")
            raise HTTPException(
                status_code=502,
                detail=f"Gotenberg conversion service error: {str(e)}"
            )

    # 2. Local LibreOffice headless binary
    soffice_bin = find_libreoffice_binary()
    if soffice_bin:
        with tempfile.TemporaryDirectory() as tmp_dir:
            cmd = [
                soffice_bin,
                "--headless",
                "--invisible",
                "--nologo",
                "--nodefault",
                "--nofirststartwizard",
                "--convert-to", "pdf",
                "--outdir", str(tmp_dir),
                str(input_path.resolve()),
            ]
            try:
                proc = subprocess.run(
                    cmd,
                    stdout=subprocess.PIPE,
                    stderr=subprocess.PIPE,
                    timeout=60,
                    check=False
                )
                if proc.returncode != 0:
                    err_msg = proc.stderr.decode("utf-8", errors="replace")
                    logger.error(f"LibreOffice conversion failed: {err_msg}")
                    raise RuntimeError(f"LibreOffice process failed (exit code {proc.returncode}): {err_msg}")

                generated_pdf = Path(tmp_dir) / f"{input_path.stem}.pdf"
                if not generated_pdf.exists() or generated_pdf.stat().st_size == 0:
                    raise RuntimeError("LibreOffice completed without producing an output PDF.")

                shutil.move(str(generated_pdf), str(output_path))
                return
            except subprocess.TimeoutExpired:
                raise HTTPException(status_code=504, detail="DOCX to PDF conversion timed out after 60 seconds.")
            except Exception as e:
                logger.error(f"LibreOffice conversion execution error: {str(e)}")
                raise HTTPException(status_code=500, detail=f"LibreOffice execution failed: {str(e)}")

    # 3. No genuine engine available in current runtime
    raise HTTPException(
        status_code=503,
        detail=(
            "Server-side DOCX-to-PDF conversion requires LibreOffice headless or Gotenberg service. "
            "Neither engine was detected in the active runtime. "
            "To enable this conversion, deploy using the Convertly Docker container (which packages LibreOffice) "
            "or configure GOTENBERG_URL in your backend environment."
        )
    )

def convert_docx_to_txt(input_path: Path, output_path: Path, options: Dict[str, Any]) -> None:
    doc = docx.Document(input_path)
    full_text = []
    for para in doc.paragraphs:
        full_text.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            full_text.append("\t".join(cell.text.strip() for cell in row.cells))
            
    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(full_text))

def convert_txt_to_docx(input_path: Path, output_path: Path, options: Dict[str, Any]) -> None:
    doc = docx.Document()
    with open(input_path, "r", encoding="utf-8", errors="replace") as f:
        for line in f:
            line_str = line.rstrip("\r\n")
            if line_str.startswith("# "):
                doc.add_heading(line_str[2:], level=1)
            elif line_str.startswith("## "):
                doc.add_heading(line_str[3:], level=2)
            else:
                doc.add_paragraph(line_str)
    doc.save(output_path)
